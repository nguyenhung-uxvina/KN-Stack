import json
import os
import subprocess
import sys

import openpyxl
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
REPO = os.path.abspath(os.path.join(SKILL, "..", "..", ".."))
FIX = os.path.join(REPO, "evals", "fixtures", "cad-workbook")
sys.path.insert(0, SKILL)

import fab_workbook as fw  # noqa: E402


@pytest.fixture()
def master_path(tmp_path):
    p = str(tmp_path / "WX-MASTER-DATA.xlsx")
    fw.init_master(p)
    return p


def test_init_master_seed(master_path):
    wb = openpyxl.load_workbook(master_path)
    for s in ["_META"] + fw.MASTER_SHEETS:
        assert s in wb.sheetnames
    ws = wb["MATERIALS"]
    header = [c.value for c in ws[1]]
    assert header[:5] == ["item_code", "item_name", "material_class", "uom", "rate"]
    rates = {r[0].value: r[4].value for r in ws.iter_rows(min_row=2)}
    assert rates["MAT-SS400"] == 18500


def test_load_extracts_and_bom():
    extracts = fw.load_extracts(FIX)
    assert [e["meta"]["part_id"] for _, e in extracts] == ["P-001", "P-002", "P-003"]
    bom = fw.load_master_bom(FIX)
    assert bom["p-001"]["qty"] == "2"
    assert bom["p-003"]["material"] == ""


def test_resolve_param():
    extracts = dict((e["meta"]["part_id"], e) for _, e in fw.load_extracts(FIX))
    bom = fw.load_master_bom(FIX)
    v, conf = fw.resolve_param(extracts["P-001"], bom.get("p-001"), "mass_kg")
    assert v == 4.2 and conf == "MED"
    v, _ = fw.resolve_param(extracts["P-001"], bom.get("p-001"), "qty")
    assert v == "2"
    v, _ = fw.resolve_param(extracts["P-003"], bom.get("p-003"), "material")
    assert v in (None, "")


def test_helpers(master_path):
    master = fw.load_master(master_path)
    assert fw.ops_from_process("laser+chấn") == ["laser", "chan"]
    assert fw.material_class("SS400", master) == "thep_tam"
    assert fw.material_code("NHÔM 5083", master) == "MAT-5083"
    assert fw.material_code("titan", master) == ""


def test_build_parts_bom(tmp_path, master_path):
    out = str(tmp_path / "FIXTURE_FAB-DB.xlsx")
    fw.main([FIX, "--master", master_path, "--project", "FIXTURE", "--out", out])
    wb = openpyxl.load_workbook(out)
    for s in ["_META", "PARTS", "BOM"] + fw.MASTER_SHEETS:
        assert s in wb.sheetnames
    parts = {r[0].value: r for r in wb["PARTS"].iter_rows(min_row=2)}
    assert parts["P-001"][6].value == 4.2          # mass_kg
    assert parts["P-001"][7].value == 2            # qty
    bom_header = [c.value for c in wb["BOM"][1]]
    assert bom_header == ["item_code", "item_name", "description", "qty_per_unit",
                          "uom", "material", "spec", "process"]


def _build(tmp_path, master_path):
    out = str(tmp_path / "FIXTURE_FAB-DB.xlsx")
    fw.main([FIX, "--master", master_path, "--project", "FIXTURE", "--out", out])
    return openpyxl.load_workbook(out)


def test_dinh_muc_formulas(tmp_path, master_path):
    wb = _build(tmp_path, master_path)
    ws = wb["DINH_MUC"]
    header = [c.value for c in ws[1]]
    assert header == ["part_id", "material_class", "item_code_vt", "qty", "mass_kg",
                      "operation", "key", "waste_pct", "vt_kg", "time_in_mins", "grade", "nc_gio"]
    rows = {(r[0].value, r[5].value): r for r in ws.iter_rows(min_row=2)}
    # P-001 process 'laser+chấn' → 2 dòng
    assert ("P-001", "laser") in rows and ("P-001", "chan") in rows
    r = rows[("P-001", "laser")]
    assert r[6].value == '=F2&"|"&B2'
    assert "VLOOKUP" in r[7].value and "WASTE_FACTORS" in r[7].value
    assert "1+H2/100" in r[8].value            # vt = mass*(1+waste%)*qty
    assert "LABOR_NORMS" in r[9].value


def test_du_toan_formulas_and_missing(tmp_path, master_path):
    wb = _build(tmp_path, master_path)
    ws = wb["DU_TOAN"]
    header = [c.value for c in ws[1]]
    assert header == ["part_id", "item_code_vt", "vt_kg", "rate_vt", "tt_vat_tu",
                      "nc_gio", "grade", "rate_nc", "tt_nhan_cong",
                      "workstation", "hour_rate", "tt_may", "tong_vnd"]
    body = list(ws.iter_rows(min_row=2))
    data, total = body[:-1], body[-1]
    r1 = data[0]
    assert "MATERIALS" in r1[3].value and "VLOOKUP" in r1[3].value
    assert "LABOR_RATES" in r1[7].value
    assert "WORKSTATIONS" in r1[10].value
    assert r1[12].value.startswith("=IFERROR(")
    assert total[0].value == "TỔNG" and total[12].value.startswith("=SUM(")
    # P-003 thiếu material → item_code_vt rỗng → rate VLOOKUP bọc IFERROR ra MISSING khi mở Excel;
    # python-side: norm_gaps báo thiếu
    gaps = fw.norm_gaps(dict((e["meta"]["part_id"], e) for _, e in fw.load_extracts(FIX))["P-003"],
                        fw.load_master_bom(FIX).get("p-003"), fw.load_master(master_path))
    assert "material" in gaps["missing"]
    assert "mass_kg" in gaps["missing"]


def test_qc_dims(tmp_path, master_path):
    wb = _build(tmp_path, master_path)
    ws = wb["QC_DIMS"]
    assert [c.value for c in ws[1]] == ["part_id", "feature", "value", "tolerance",
                                        "confidence", "source", "gauge"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    r1 = [r for r in rows if r[0] == "P-001"][0]
    assert r1[3] == 0.1 and r1[6] == "panme"


def test_checklist_flags_missing(tmp_path, master_path):
    wb = _build(tmp_path, master_path)
    ws = wb["CHECKLIST"]
    assert [c.value for c in ws[1]] == ["part_id", "QTCN", "DINH_MUC", "SO_TAY_QC", "BOM", "DU_TOAN"]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert rows["P-001"][4].value == "ĐỦ"                       # BOM đủ
    assert "THIẾU" in rows["P-003"][4].value and "material" in rows["P-003"][4].value
    assert rows["P-003"][4].fill.fgColor.rgb.endswith("FFC7CE")  # đỏ
    assert str(rows["P-002"][1].value).startswith("ĐỦ")   # holes/surface chỉ là warning phụ


def test_check_part_api(master_path):
    master = fw.load_master(master_path)
    reqs = json.load(open(os.path.join(SKILL, "references", "param_requirements.json"), encoding="utf-8"))
    e3 = dict((e["meta"]["part_id"], e) for _, e in fw.load_extracts(FIX))["P-003"]
    res = fw.check_part(e3, fw.load_master_bom(FIX).get("p-003"), master, reqs)
    assert "material" in res["BOM"]["missing_critical"]
    assert "mass_kg" in res["DU_TOAN"]["missing_critical"]
    assert res["SO_TAY_QC"]["missing_warning"]  # tolerances/holes trống


def test_refresh_creates_bak(tmp_path, master_path):
    out = str(tmp_path / "FIXTURE_FAB-DB.xlsx")
    fw.main([FIX, "--master", master_path, "--project", "FIXTURE", "--out", out])
    fw.main([FIX, "--master", master_path, "--project", "FIXTURE", "--out", out])
    assert os.path.exists(out + ".bak")


def test_resolve_param_thickness_multirow_bom():
    # parity với validate.py _thickness_rows: bom[0] thiếu thickness nhưng bom[1] có
    e = {"meta": {"part_id": "X"}, "bom": [{"code": "X"}, {"code": "X2", "thickness_mm": 5.0}],
         "dimensions": [], "tolerances": [], "holes": [], "conflicts": [], "process_notes": []}
    v, conf = fw.resolve_param(e, None, "thickness_mm")
    assert v == 5.0 and conf == "MED"
