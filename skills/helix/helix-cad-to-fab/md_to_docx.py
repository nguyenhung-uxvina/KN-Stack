#!/usr/bin/env python3
"""
md_to_docx.py — Markdown → DOCX for Vietnamese defense process documents.

Handles: ATX headings (#..####), GFM tables (incl. aligned/padded), bullet
lists, **bold**, blockquotes (>), horizontal rules. Times New Roman 11pt.
Falls back to a *_v2.docx name if the target is locked open in Word.

  python md_to_docx.py <in.md> [out.docx]

Run with PYTHONUTF8=1 on Windows.
"""
import os
import re
import sys


def convert(src, dst=None):
    from docx import Document
    from docx.shared import Pt
    dst = dst or os.path.splitext(src)[0] + ".docx"
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)

    def add_runs(p, text):
        for seg in re.split(r"(\*\*.+?\*\*)", text):
            if not seg:
                continue
            if seg.startswith("**") and seg.endswith("**"):
                r = p.add_run(seg[2:-2]); r.bold = True
            else:
                p.add_run(seg)

    i, n = 0, len(lines)
    while i < n:
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.lstrip().startswith("|"):
            tbl = []
            while i < n and lines[i].lstrip().startswith("|"):
                tbl.append(lines[i]); i += 1
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in tbl]
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
            if rows:
                t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                t.style = "Light Grid Accent 1"
                for ri, r in enumerate(rows):
                    for ci, c in enumerate(r):
                        cell = t.rows[ri].cells[ci]; cell.text = ""
                        pr = cell.paragraphs[0]; add_runs(pr, c.replace("**", ""))
                        if ri == 0:
                            for rr in pr.runs:
                                rr.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1))); i += 1; continue
        if ln.strip() == "---":
            doc.add_paragraph("_" * 40); i += 1; continue
        if ln.lstrip().startswith(">"):
            p = doc.add_paragraph(); p.style = "Intense Quote"
            add_runs(p, ln.lstrip()[1:].strip()); i += 1; continue
        if re.match(r"^\s*[-*]\s+", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*[-*]\s+", "", ln)); i += 1; continue
        p = doc.add_paragraph(); add_runs(p, ln); i += 1

    for cand in (dst, os.path.splitext(dst)[0] + "_v2.docx"):
        try:
            doc.save(cand)
            return cand
        except PermissionError:
            continue
    raise PermissionError(f"{dst} locked (đang mở trong Word)")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: python md_to_docx.py <in.md> [out.docx]")
    out = convert(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    print("wrote", out, os.path.getsize(out), "bytes")
